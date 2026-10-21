#!/usr/bin/env python3
'Map [PRIVATE_SUBJECT] age, sex, send date, and material documentation to patient keys.\n\nThe original workbook subject identifier is used only in memory for an exact join.\nrecorded_n_binary is carried for auditing and is never a predictor.\n'
from __future__ import annotations
from release_runtime import private_literal as _private_literal, release_path as _release_path

import argparse
import csv
import hashlib
import io
import json
import math
import os
from collections import Counter
from datetime import date, datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(_release_path('path_1304218e86d2e596'))
SOURCE = Path(_release_path('path_df0fb75e17e2f0ab'))
CLINICAL = SOURCE / "ln_CRC_TCR_Clinical_data.xlsx"
LIBRARY = SOURCE / "TCR-H2498建库报告.xlsx"
LEDGER = Path(_release_path('path_67211e466e7fa9cd'))
COHORT = ROOT / "config/cohort_samples.tsv"
BLOOD31 = ROOT / "config/documented_blood31.tsv"
OUT = ROOT / "config/home_clinical_covariates.tsv"
AUDIT = ROOT / "audit/00_clinical_baseline_provenance.json"
PINS = {
    "clinical_workbook": (CLINICAL, "7e1529063bf60fccdde92906782887f3a7cd2e9537892f8c94f45db17490dc84"),
    "library_workbook": (LIBRARY, "f805d320fddf1ce30c4a4bc14d92b8cc500794cf7f0f0180b36479d9ccdf9d77"),
    "home_ledger": (LEDGER, "32d80d41cd81ccf6f077aeca607a757641f31a4756974a8be4f71f1ba54f7c1d"),
}
FIELDS = [
    "patient_key", "age_years", "sex", "library_batch", "group_key", "sample_id", "send_date",
    "material_verified_blood", "material_status", "identity_provenance",
    "clinical_sheet_row", "library_sheet_row", "recorded_n_binary",
]
FEATURE_COLUMNS = ["age_years", "sex", "library_batch"]
AUDIT_ONLY_COLUMNS = [
    "patient_key", "group_key", "sample_id", "material_verified_blood",
    "material_status", "identity_provenance", "clinical_sheet_row",
    "library_sheet_row", "recorded_n_binary",
]
BLOOD_TEXT = "结直肠癌患者外周血"


def require(value: bool, message: str) -> None:
    if not value:
        raise RuntimeError("CLINICAL_BASELINE_STOP: " + message)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(8 << 20), b""):
            h.update(block)
    return h.hexdigest()


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f, delimiter="\t"))


def text(value: object) -> str:
    return "" if value is None else str(value).strip()


def iso_date(value: object) -> str:
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    candidate = text(value)
    date.fromisoformat(candidate)
    return candidate


def make_bytes(rows: list[dict[str, str]]) -> bytes:
    buf = io.StringIO(newline="")
    writer = csv.DictWriter(buf, fieldnames=FIELDS, delimiter="\t", lineterminator="\n")
    writer.writeheader()
    writer.writerows(rows)
    return buf.getvalue().encode("utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--verify", action="store_true",
                        help="recompute and compare the existing output without writing")
    args = parser.parse_args()

    hashes: dict[str, str] = {}
    for name, (path, expected) in PINS.items():
        hashes[name] = sha256(path)
        require(hashes[name] == expected, f"{name} SHA256 differs from frozen source")

    clinical_sheet = load_workbook(CLINICAL, read_only=True, data_only=True)["Sheet1"]
    library_sheet = load_workbook(LIBRARY, read_only=True, data_only=True)["TCR-BCR建库数据"]
    clinical = list(clinical_sheet.values)
    library = list(library_sheet.values)
    require(len(clinical) == 140 and len(library) == 139, "source workbook row count")
    require(text(clinical[0][9]) == "病理报告" and
            [text(clinical[1][i]) for i in (0, 2, 3, 6, 11)] ==
            ["样品名称", "age", "sex", "Neoadjuvant thearpy", "N"],
            "clinical workbook headers")
    require([text(library[0][i]) for i in (1, 2, 3, 4, 6)] ==
            ["发送日期", "康测编号", "样品名称", "样品类别", "建库结果"],
            "library workbook headers")
    clinical_ids = [text(r[0]) for r in clinical[2:]]
    library_ids = [text(r[3]) for r in library[1:]]
    sample_ids = [text(r[2]) for r in library[1:]]
    require(len(set(clinical_ids)) == 138 and set(clinical_ids) == set(library_ids)
            and len(set(sample_ids)) == 138 and all(clinical_ids + library_ids + sample_ids),
            "source subject/sample identifiers are not exact one-to-one")

    ledger = read_tsv(LEDGER)
    chosen = [r for r in ledger if r[_private_literal('private_b83fec3a2cefc253')] == "yes"]
    require(len(ledger) == 138 and len(chosen) == 114, "Home ledger population")
    require(len({r["patient_key"] for r in chosen}) == 114 and
            len({r["group_key"] for r in chosen}) == 114 and
            len({r["sample_id"] for r in chosen}) == 114, _private_literal('private_12fb2cc30d7e422d'))

    cohort = [r for r in read_tsv(COHORT) if r["cohort"] == "home"]
    require(len(cohort) == 114 and
            {(r["sample_id"], r["patient_key"], r["recorded_n_binary"]) for r in cohort} ==
            {(r["sample_id"], r["patient_key"], r["pathology_report_N_binary"]) for r in chosen},
            _private_literal('private_0e67846c711df55d'))

    out_rows: list[dict[str, str]] = []
    for row in sorted(chosen, key=lambda x: x["patient_key"]):
        ci = int(row["clinical_sheet_row"]) - 1
        li = int(row["library_sheet_row"]) - 1
        require(2 <= ci < len(clinical) and 1 <= li < len(library),
                "source row number out of range")
        cr, lr = clinical[ci], library[li]
        require(text(lr[2]) == row["sample_id"] and text(cr[0]) == text(lr[3]),
                "sample or original subject join mismatch")
        age_raw = cr[2]
        require(isinstance(age_raw, (int, float)) and not isinstance(age_raw, bool) and
                math.isfinite(float(age_raw)) and float(age_raw).is_integer() and
                0 < age_raw < 120, "age missing/invalid")
        age = str(int(age_raw))
        sex = text(cr[3]).lower()
        require(sex in {"f", "m"}, "sex missing/invalid")
        treatment = text(cr[6]).lower()
        require(treatment in {"n", "no"} and row["neoadjuvant_normalized"] == "no",
                _private_literal('private_ecaeffbd34504fe1'))
        require(text(cr[11]) == row["pathology_report_N_raw"] and
                row["pathology_report_N_binary"] in {"0", "1"},
                "pathology report N join mismatch")
        require(text(lr[6]) == "建库成功", _private_literal('private_b9305672b3cef9d9'))
        send_date = iso_date(lr[1])
        batch = row["sample_id"][:6]
        require(batch in {"H2498B", "H2498C", "H2498D"}, "unexpected sample batch")
        material = text(lr[4])
        blood = material == BLOOD_TEXT
        require((row["explicit_peripheral_blood_in_report"] == "yes") == blood,
                "material documentation differs from frozen ledger")
        out_rows.append({
            "patient_key": row["patient_key"],
            "group_key": row["group_key"],
            "sample_id": row["sample_id"],
            "age_years": age,
            "sex": sex,
            "library_batch": batch,
            "send_date": send_date,
            "material_verified_blood": "yes" if blood else "unverified",
            "material_status": "explicit_peripheral_blood" if blood else "unverified_material",
            "identity_provenance": "exact_source_subject_id_join",
            "clinical_sheet_row": row["clinical_sheet_row"],
            "library_sheet_row": row["library_sheet_row"],
            "recorded_n_binary": row["pathology_report_N_binary"],
        })

    by_batch = {b: {r["send_date"] for r in out_rows if r["library_batch"] == b}
                for b in ("H2498B", "H2498C", "H2498D")}
    require(by_batch == {"H2498B": {"2024-08-21"}, "H2498C": {"2024-10-11"},
                         "H2498D": {"2024-11-22"}}, "batch/date mapping changed")
    blood_rows = [r for r in out_rows if r["material_verified_blood"] == "yes"]
    require(len(blood_rows) == 31 and Counter(r["recorded_n_binary"] for r in blood_rows) ==
            {"0": 13, "1": 18}, "documented blood subset changed")
    published_blood = read_tsv(BLOOD31)
    require({(r["patient_key"], r["group_key"], r["sample_id"], r["recorded_n_binary"])
             for r in published_blood} ==
            {(r["patient_key"], r["group_key"], r["sample_id"], r["recorded_n_binary"])
             for r in blood_rows}, "documented blood31 table differs from source")
    hashes["cohort_samples"] = sha256(COHORT)
    hashes["documented_blood31"] = sha256(BLOOD31)
    result = make_bytes(out_rows)
    output_hash = hashlib.sha256(result).hexdigest()

    if args.verify:
        require(OUT.is_file() and OUT.read_bytes() == result, "saved covariate table differs")
        receipt = json.loads(AUDIT.read_text(encoding="utf-8"))
        require(receipt["output_sha256"] == output_hash and
                receipt["source_sha256"] == hashes and receipt["status"] == "PASS",
                "saved provenance receipt differs")
        print(json.dumps({"status": "PASS_VERIFY", "n": len(out_rows),
                          "output_sha256": output_hash}, sort_keys=True))
        return

    require(not OUT.exists() and not AUDIT.exists(), "output already exists; use --verify")
    OUT.write_bytes(result)
    receipt = {
        "schema": _private_literal('private_696b762cd969621d'),
        "status": "PASS",
        "created_utc": datetime.now(timezone.utc).isoformat(),
        "source_sha256": hashes,
        "output_path": str(OUT),
        "output_sha256": output_hash,
        "n_home": len(out_rows),
        "unique_patient_keys": len({r["patient_key"] for r in out_rows}),
        "recorded_n_counts": dict(Counter(r["recorded_n_binary"] for r in out_rows)),
        "documented_blood_counts": dict(Counter(r["material_verified_blood"] for r in out_rows)),
        "batch_counts": dict(Counter(r["library_batch"] for r in out_rows)),
        "missing_age_sex_send_date": {
            k: sum(not r[k] for r in out_rows) for k in ("age_years", "sex", "send_date")
        },
        "identity_join": "library sample ID to frozen ledger; original subject ID exactly equals clinical workbook subject ID, used in memory only",
        "feature_columns": FEATURE_COLUMNS,
        "audit_only_columns": AUDIT_ONLY_COLUMNS,
        "baseline_policy": {
            "clinical_primary": ["age_years", "sex"],
            "technical": ["library_batch", "log1p_productive_reads", "log1p_productive_unique_aa"],
            "do_not_fit_date_and_batch_together": True,
            "exclude_from_predictors": ["recorded_n_binary", "patient_key", "group_key",
                                        "sample_id", "clinical_sheet_row", "library_sheet_row",
                                        "identity_provenance", "material_verified_blood",
                                        "material_status"],
            "treatment_note": _private_literal('private_26894a82bbfbca50'),
            "pathology_note": "T/N/M, stage Group, grade, MSI/MMR, EMVI, and PNI are not baseline predictors",
            "blood_note": "Only exact library text documents blood for 31; 83 remain unverified",
        },
    }
    AUDIT.write_text(json.dumps(receipt, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
                     encoding="utf-8")
    print(json.dumps({"status": "PASS", "n": len(out_rows),
                      "output_sha256": output_hash}, sort_keys=True))


if __name__ == "__main__":
    main()
